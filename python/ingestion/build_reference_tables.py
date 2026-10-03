"""
Build Reference Lookup Tables from the Official DfT STATS19 Data Guide.
Extracts code-to-label mappings into standardized CSVs in data/reference/.
"""

import sys
from pathlib import Path
import openpyxl
import pandas as pd

# Add repo root to sys.path
SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent.parent
sys.path.insert(0, str(REPO_ROOT))

import python.config as config

GUIDE_FILE_PATH = config.REFERENCE_DATA_DIR / "dft_road_safety_data_guide.xlsx"


def build_all_reference_tables():
    """Extract code-to-label tables from the official Excel data guide."""
    print("=" * 70)
    print("Building Reference Lookup Tables from STATS19 Data Guide")
    print("=" * 70)

    if not GUIDE_FILE_PATH.exists():
        print(f"[ERROR] Data guide not found at: {GUIDE_FILE_PATH}")
        print("Please run python/ingestion/download_data.py first.")
        return False

    config.REFERENCE_DATA_DIR.mkdir(parents=True, exist_ok=True)

    # Load data guide
    print(f"Loading data guide: {GUIDE_FILE_PATH.name}...")
    wb = openpyxl.load_workbook(GUIDE_FILE_PATH, data_only=True)

    if "2024_code_list" not in wb.sheetnames:
        print(f"[ERROR] '2024_code_list' sheet not found in {wb.sheetnames}")
        return False

    sheet = wb["2024_code_list"]
    raw_data = [row for row in sheet.iter_rows(values_only=True)]
    df_codes = pd.DataFrame(raw_data[1:], columns=raw_data[0])

    # Clean column names
    df_codes.columns = [str(c).strip().lower().replace(" ", "_").replace("/", "_") for c in df_codes.columns]

    # Filter rows that have valid code/format and label
    df_valid = df_codes.dropna(subset=["code_format", "label"]).copy()
    df_valid["code_format"] = df_valid["code_format"].astype(str).str.strip()
    df_valid["label"] = df_valid["label"].astype(str).str.strip()

    # Save complete raw lookup reference
    all_lookups_csv = config.REFERENCE_DATA_DIR / "all_reference_codes.csv"
    df_valid.to_csv(all_lookups_csv, index=False)
    print(f"Saved master reference table: {all_lookups_csv.name} ({len(df_valid):,} entries)")

    # Extract individual categorical domain tables
    fields = df_valid["field_name"].dropna().unique()
    extracted_count = 0

    print("\nExtracting categorical reference CSVs:")
    for field in sorted(fields):
        sub_df = df_valid[df_valid["field_name"] == field][["code_format", "label", "note", "table"]].copy()
        sub_df = sub_df.rename(columns={"code_format": "code"})
        
        # Deduplicate on code
        sub_df = sub_df.drop_duplicates(subset=["code"])

        out_path = config.REFERENCE_DATA_DIR / f"{field}.csv"
        sub_df.to_csv(out_path, index=False)
        extracted_count += 1
        print(f"  -> {out_path.name} ({len(sub_df)} codes)")

    # Ensure critical target and feature domains exist with verified fallbacks
    ensure_domain_tables()

    print("\n" + "=" * 70)
    print(f"[SUCCESS] Extracted {extracted_count} reference lookup tables into {config.REFERENCE_DATA_DIR}")
    print("=" * 70)
    return True


def ensure_domain_tables():
    """Ensure standard explicit reference files exist with clean schemas."""
    # Day of week reference
    dow_path = config.REFERENCE_DATA_DIR / "day_of_week.csv"
    if not dow_path.exists():
        dow_df = pd.DataFrame([
            {"code": 1, "label": "Sunday"},
            {"code": 2, "label": "Monday"},
            {"code": 3, "label": "Tuesday"},
            {"code": 4, "label": "Wednesday"},
            {"code": 5, "label": "Thursday"},
            {"code": 6, "label": "Friday"},
            {"code": 7, "label": "Saturday"},
        ])
        dow_df.to_csv(dow_path, index=False)

    # Time period proxy reference
    tp_path = config.REFERENCE_DATA_DIR / "time_period.csv"
    tp_df = pd.DataFrame([
        {"code": 1, "label": "Morning Peak", "hours": "07:00-09:59", "description": "Morning rush hour commuter traffic proxy"},
        {"code": 2, "label": "Daytime Off-Peak", "hours": "10:00-15:59", "description": "Daytime inter-peak traffic proxy"},
        {"code": 3, "label": "Evening Peak", "hours": "16:00-18:59", "description": "Evening rush hour commuter traffic proxy"},
        {"code": 4, "label": "Night / Off-Hours", "hours": "19:00-06:59", "description": "Nighttime and early morning low traffic proxy"}
    ])
    tp_df.to_csv(tp_path, index=False)

    # Speed limit reference
    speed_path = config.REFERENCE_DATA_DIR / "speed_limit.csv"
    if not speed_path.exists():
        speed_df = pd.DataFrame([
            {"code": 20, "label": "20 mph"},
            {"code": 30, "label": "30 mph"},
            {"code": 40, "label": "40 mph"},
            {"code": 50, "label": "50 mph"},
            {"code": 60, "label": "60 mph"},
            {"code": 70, "label": "70 mph"},
            {"code": -1, "label": "Data missing or out of range"},
            {"code": 99, "label": "Unknown"}
        ])
        speed_df.to_csv(speed_path, index=False)

    # Add code 33 (Powered personal transporter / e-scooter) to casualty_type if not present
    ct_path = config.REFERENCE_DATA_DIR / "casualty_type.csv"
    if ct_path.exists():
        ct_df = pd.read_csv(ct_path)
        if "33" not in ct_df["code"].astype(str).values:
            new_row = pd.DataFrame([{"code": 33, "label": "Powered personal transporter", "note": "introduced in recent DfT specification", "table": "casualty"}])
            ct_df = pd.concat([ct_df, new_row], ignore_index=True)
            ct_df.to_csv(ct_path, index=False)


if __name__ == "__main__":
    success = build_all_reference_tables()
    if not success:
        sys.exit(1)
